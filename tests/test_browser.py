"""Real-browser checks for the integrated public entry, run with pytest -m browser."""
from io import BytesIO
from pathlib import Path
import shutil
import signal
import socket
import subprocess
import sys
import time
from urllib.request import urlopen

from openpyxl import Workbook, load_workbook
import pytest
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
pytestmark = pytest.mark.browser


@pytest.fixture(scope='module')
def site():
    with socket.socket() as sock:
        sock.bind(('127.0.0.1', 0))
        port = sock.getsockname()[1]
    process = subprocess.Popen([
        sys.executable, '-m', 'streamlit', 'run', 'streamlit_app.py',
        '--server.headless', 'true', '--server.address', '127.0.0.1',
        '--server.port', str(port), '--browser.gatherUsageStats', 'false',
    ], cwd=ROOT, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    url = f'http://127.0.0.1:{port}'
    try:
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            if process.poll() is not None:
                pytest.fail('Streamlit exited before readiness')
            try:
                with urlopen(url + '/_stcore/health', timeout=1) as response:
                    if response.status == 200:
                        break
            except OSError:
                time.sleep(0.1)
        else:
            pytest.fail('Streamlit did not become healthy')
        yield url
    finally:
        process.send_signal(signal.SIGINT)
        try:
            process.wait(timeout=10)
        except subprocess.TimeoutExpired:
            process.kill()
            process.wait(timeout=5)


@pytest.mark.parametrize('viewport', [{'width': 1440, 'height': 1000}, {'width': 390, 'height': 844}], ids=['desktop', 'mobile'])
def test_integrated_search_import_download_random_and_brand(site, viewport, tmp_path):
    executable = shutil.which('chromium')
    if not executable:
        pytest.skip('Install Chromium to run these browser checks')
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(executable_path=executable, headless=True, args=['--no-sandbox'])
        page = browser.new_page(viewport=viewport)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        try:
            page.goto(site, wait_until='domcontentloaded')
            expect(page.locator('.card-heading')).to_have_count(14, timeout=30000)
            expect(page).to_have_title('清食指南｜清大餐廳選擇器')
            page.get_by_text('搜尋與抽選', exact=True).click()
            expect(page.locator('.wordmark')).to_have_text('清食指南')
            expect(page.locator('.st-key-catalog_grid .food-card')).to_have_count(6)
            expect(page.get_by_text('評論尚未啟用', exact=False)).to_be_visible()
            query = page.get_by_role('textbox', name='搜尋店家', exact=True)
            query.fill('麥當牢')
            results = page.locator('.st-key-search_results')
            expect(results).to_contain_text('麥當勞', timeout=15000)
            expect(results).to_contain_text('讀音或字詞相近')
            expect(query).to_be_focused()
            page.get_by_role('button', name='清除搜尋', exact=True).click()
            expect(query).to_have_value('')
            page.get_by_role('button', name='下一頁', exact=True).click()
            expect(page.get_by_text('第 2 / 3 頁', exact=True)).to_be_visible()
            page.get_by_role('button', name='下一頁', exact=True).click()
            expect(page.locator('.st-key-catalog_grid .food-card')).to_have_count(2)
            page.get_by_role('button', name='幫我決定這一餐', exact=True).click()
            picked = page.locator('[aria-label="餐廳推薦結果"]')
            expect(picked).to_be_visible()
            first = picked.locator('h2').inner_text()
            page.get_by_role('button', name='再推薦一家', exact=True).click()
            expect(picked.locator('h2')).not_to_have_text(first)
            workbook = Workbook()
            workbook.active.append(['Restaurant', 'Location', 'Gener', 'Price'])
            workbook.active.append(['匯入瀏覽器測試店', '小吃部', '早餐', 1.5])
            upload = tmp_path / 'custom.xlsx'
            workbook.save(upload)
            workbook.close()
            page.get_by_role('button', name='匯入 Excel 名單', exact=True).click()
            dialog = page.get_by_role('dialog')
            expect(dialog).to_be_visible()
            dialog.locator('input[type=file]').set_input_files(upload)
            dialog.get_by_role('button', name='確認匯入', exact=True).click()
            expect(page.locator('.st-key-catalog_grid')).to_contain_text('匯入瀏覽器測試店')
            expect(page.locator('.st-key-catalog_grid .price-half-coin')).to_have_count(1)
            half = page.locator('.st-key-catalog_grid .price-half-coin')
            assert half.bounding_box()['width'] < half.locator('span').bounding_box()['width']
            link = page.locator('.st-key-catalog_grid .poster-link')
            assert 'google.com/maps/search/' in link.get_attribute('href')
            page.get_by_text('查看全部符合條件的餐廳（1 間）', exact=True).click()
            with page.expect_download() as event:
                page.get_by_role('button', name='匯出篩選名單', exact=True).click()
            book = load_workbook(BytesIO(Path(event.value.path()).read_bytes()), read_only=True)
            try:
                values = list(book.active.values)
                assert len(values) == 2
                assert values[1][0] == '匯入瀏覽器測試店' and values[1][3] == 1.5
            finally:
                book.close()
            page.get_by_role('button', name='恢復內建名單', exact=True).click()
            expect(page.locator('.st-key-catalog_grid .food-card')).to_have_count(6)
            expect(page.get_by_text('第 1 / 3 頁', exact=True)).to_be_visible()
            assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
            assert not errors, errors
        finally:
            browser.close()
